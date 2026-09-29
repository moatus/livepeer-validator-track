"""Independent fixture checks and inspection of the saved XLSX formulas/caches."""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import load_workbook

from build import OUTPUT, UNDEFINED, account


def close(actual, expected, label):
    if isinstance(expected, str):
        assert actual == expected, f"{label}: {actual!r} != {expected!r}"
    else:
        assert isinstance(actual, (int, float)) and math.isclose(actual, expected, abs_tol=1e-9, rel_tol=1e-9), f"{label}: {actual!r} != {expected!r}"


def fixtures():
    # These constants are hand-computable from 940 LPT node budget, not read from the workbook.
    cases = [
        ("aligned 60/40", (600, 400), (600, 400), (1, 1), 940, 0, 0, (564, 376)),
        ("mismatch 90/10 vs 10/90", (900, 100), (100, 900), (1, 1), 188, 752, 0, (94, 94)),
        ("A score reduced", (600, 400), (600, 400), (.5, 1), 658, 0, 282, (282, 376)),
        ("A no stake, positive network stake", (600, 400), (0, 1000), (1, 1), 376, 564, 0, (0, 376)),
    ]
    for name, fees, stakes, scores, paid, gap, withheld, each in cases:
        result = account(fees, stakes, scores)
        close(result["paid"], paid, name+" paid")
        close(result["unallocated"], gap, name+" MFS gap")
        close(result["withheld"], withheld, name+" score gap")
        for i in range(2): close(result["rows"][i]["paid"], each[i], name+f" node {i}")
    no_stake = account((600, 400), (0, 1000), (1, 1))
    close(no_stake["rows"][0]["node_usdc"], 300, "no-stake ordinary fee")
    for fees, stakes in (((0, 0), (600, 400)), ((600, 400), (0, 0)), ((0, 0), (0, 0))):
        result = account(fees, stakes, (1, 1))
        assert result["paid"] is None and result["unallocated"] is None and result["withheld"] is None
        assert all(r["mfs"] is None and r["paid"] is None for r in result["rows"])
    print("7 independent arithmetic fixtures: PASS")


def workbook():
    assert OUTPUT.is_file(), f"Build first: {OUTPUT}"
    vals = load_workbook(OUTPUT, data_only=True)
    forms = load_workbook(OUTPUT, data_only=False)
    assert vals.sheetnames == ["Inputs", "Cooperative", "Ledger & limits", "One-variable", "Pass-through margin", "Source & decisions"]
    formula_count = 0
    for sheet in forms.sheetnames:
        for row in forms[sheet]:
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
                    assert vals[sheet][cell.coordinate].value is not None, f"Missing cache: {sheet}!{cell.coordinate}"
                assert vals[sheet][cell.coordinate].data_type != "e", f"Excel error: {sheet}!{cell.coordinate}"
    assert formula_count >= 190, formula_count
    for addr in ("B16", "B17", "B18"): close(vals["Inputs"][addr].value, "OK", "Inputs "+addr)
    for addr, n in {"B6":600,"C6":300,"D6":240,"E6":60,"F6":300,"H6":60,"I6":.6,"J6":.6,"K6":.6,"M6":564,"N6":564,"O6":169.2,"P6":394.8,
                    "B7":400,"C7":200,"D7":160,"E7":40,"F7":200,"H7":30,"I7":.4,"J7":.4,"K7":.4,"M7":376,"N7":376,"O7":112.8,"P7":263.2,
                    "B9":1000,"C9":500,"D9":400,"E9":100,"F9":500,"H9":90,"M9":940,"N9":940,"O9":282,"P9":658}.items():
        close(vals["Cooperative"][addr].value, n, "Cooperative "+addr)
    for addr, n in {"B6":1000,"B7":500,"B8":500,"B9":400,"B10":100,"B11":90,"B13":250,"B15":1000,"B16":940,"B17":10,"B18":50,"B20":0,"B21":0,"B22":940,"B23":282,"B24":658,"B26":"OK","B27":"OK"}.items():
        close(vals["Ledger & limits"][addr].value, n, "Ledger "+addr)
    # Each sensitivity cache is compared to a fresh calculation with the other node fixed.
    for row in range(7, 14):
        a_fee = vals["One-variable"][f"A{row}"].value
        m = account((a_fee, 400), (600, 400), (1, 1))
        expected = {"B":m["fees"],"C":m["rows"][0]["fee_share"],"D":m["rows"][0]["stake_share"],
                    "E":m["rows"][0]["mfs"],"F":m["rows"][1]["mfs"],"G":sum(x["mfs"] for x in m["rows"]),
                    "H":m["rows"][0]["paid"],"I":m["rows"][1]["paid"],"J":m["unallocated"],"K":m["withheld"]}
        for col, value in expected.items(): close(vals["One-variable"][f"{col}{row}"].value, value, f"fee sensitivity {col}{row}")
    for row in range(18, 25):
        a_stake = vals["One-variable"][f"A{row}"].value
        m = account((600, 400), (a_stake, 400), (1, 1))
        expected = {"B":m["stake"],"C":m["rows"][0]["fee_share"],"D":m["rows"][0]["stake_share"],
                    "E":m["rows"][0]["mfs"],"F":m["rows"][1]["mfs"],"G":sum(x["mfs"] for x in m["rows"]),
                    "H":m["rows"][0]["paid"],"I":m["rows"][1]["paid"],"J":m["unallocated"],"K":m["withheld"]}
        for col, value in expected.items(): close(vals["One-variable"][f"{col}{row}"].value, value, f"stake sensitivity {col}{row}")
    close(vals["Pass-through margin"]["B12"].value, .4, "pass-through receipt")
    close(vals["Pass-through margin"]["B13"].value, -.02, "pass-through margin")
    close(vals["Pass-through margin"]["B14"].value, 1.05, "pass-through break-even")
    formula_expect = {
        ("Cooperative","K6"): "MIN(I6,J6)",
        ("Cooperative","M6"): "Inputs!$B$14*Inputs!$B$8*K6",
        ("Cooperative","N6"): "M6*L6",
        ("Cooperative","M9"): "COUNT(M6:M7)=2",
        ("Cooperative","N9"): "COUNT(N6:N7)=2",
        ("Ledger & limits","B20"): "B16-Cooperative!M9",
        ("Ledger & limits","B21"): "Cooperative!M9-Cooperative!N9",
        ("One-variable","C10"): "A10/B10",
        ("One-variable","D22"): "A22/B22",
    }
    for (sheet, cell), fragment in formula_expect.items():
        assert fragment in forms[sheet][cell].value, (sheet, cell, forms[sheet][cell].value)
    assert UNDEFINED in forms["Cooperative"]["I6"].value
    assert UNDEFINED in forms["Cooperative"]["J6"].value
    assert UNDEFINED in forms["One-variable"]["D18"].value
    print(f"{formula_count} saved formulas, all cached; default and 14 sensitivity rows: PASS")


if __name__ == "__main__":
    fixtures()
    workbook()
