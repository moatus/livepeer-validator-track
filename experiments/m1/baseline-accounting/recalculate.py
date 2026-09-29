"""Optional LibreOffice regression of edited workbook copies; never changes the deliverable."""

from __future__ import annotations

import hashlib
import math
import shutil
import subprocess
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from build import CHECK, OUTPUT, UNDEFINED


HERE = Path(__file__).resolve().parent
NO_FINITE = "no finite break-even"


def expect(actual, target, label):
    if isinstance(target, str):
        assert actual == target, f"{label}: {actual!r} != {target!r}"
    else:
        assert isinstance(actual, (int, float)) and math.isclose(actual, target, abs_tol=1e-9, rel_tol=1e-9), f"{label}: {actual!r} != {target!r}"


def case(name, edits, expected):
    return name, edits, expected


CASES = [
    case("text_fee", {("Inputs", "B23"): "oops"}, {("Inputs", "B18"): CHECK, ("Cooperative", "C6"): CHECK, ("Cooperative", "I6"): CHECK, ("Cooperative", "B9"): CHECK, ("Ledger & limits", "B26"): CHECK, ("Ledger & limits", "B27"): CHECK, ("One-variable", "H10"): CHECK}),
    case("blank_fee", {("Inputs", "B23"): None}, {("Inputs", "B18"): CHECK, ("Cooperative", "B6"): CHECK, ("Cooperative", "B9"): CHECK, ("Ledger & limits", "B26"): CHECK}),
    case("negative_fee", {("Inputs", "B23"): -1}, {("Inputs", "B18"): CHECK, ("Cooperative", "N6"): CHECK, ("Ledger & limits", "B27"): CHECK}),
    case("nan_text", {("Inputs", "B23"): "NaN"}, {("Inputs", "B18"): CHECK, ("Cooperative", "C6"): CHECK}),
    case("error_fee", {("Inputs", "B23"): "=1/0"}, {("Inputs", "B18"): CHECK, ("Cooperative", "C6"): CHECK}),
    case("blank_budget", {("Inputs", "B14"): None}, {("Inputs", "B18"): CHECK, ("Ledger & limits", "B15"): CHECK, ("One-variable", "H10"): CHECK}),
    case("negative_buyback_price", {("Inputs", "B12"): -2}, {("Inputs", "B18"): CHECK, ("Ledger & limits", "B13"): CHECK}),
    case("blank_operator_cost", {("Inputs", "G23"): None}, {("Inputs", "B18"): CHECK, ("Cooperative", "H6"): CHECK, ("Cooperative", "H9"): CHECK}),
    case("text_retention", {("Inputs", "E23"): "oops"}, {("Inputs", "B18"): CHECK, ("Cooperative", "D6"): CHECK, ("Pass-through margin", "B11"): CHECK}),
    case("invalid_fee_split", {("Inputs", "B5"): .6}, {("Inputs", "B16"): CHECK, ("Inputs", "B18"): CHECK, ("Ledger & limits", "B26"): CHECK}),
    case("text_fee_split", {("Inputs", "B5"): "oops"}, {("Inputs", "B16"): CHECK, ("Inputs", "B18"): CHECK, ("Cooperative", "B9"): CHECK}),
    case("invalid_emission_share", {("Inputs", "B8"): .95}, {("Inputs", "B17"): CHECK, ("Inputs", "B18"): CHECK, ("One-variable", "K10"): CHECK}),
    case("text_emission_share", {("Inputs", "B8"): "oops"}, {("Inputs", "B17"): CHECK, ("Inputs", "B18"): CHECK, ("Ledger & limits", "B16"): CHECK}),
    case("invalid_score", {("Inputs", "D23"): 1.2}, {("Inputs", "B18"): CHECK, ("Cooperative", "N6"): CHECK, ("One-variable", "L10"): CHECK, ("One-variable", "H10"): CHECK, ("One-variable", "K10"): CHECK}),
    case("negative_scenario_fee", {("One-variable", "A7"): -100}, {("Inputs", "B18"): "OK", ("One-variable", "L7"): CHECK, ("One-variable", "B7"): CHECK, ("One-variable", "H7"): CHECK, ("One-variable", "K7"): CHECK}),
    case("text_scenario_fee", {("One-variable", "A7"): "oops"}, {("One-variable", "L7"): CHECK, ("One-variable", "B7"): CHECK, ("One-variable", "H7"): CHECK}),
    case("blank_scenario_stake", {("One-variable", "A18"): None}, {("One-variable", "L18"): CHECK, ("One-variable", "B18"): CHECK, ("One-variable", "H18"): CHECK}),
    case("negative_scenario_stake", {("One-variable", "A18"): -100}, {("One-variable", "L18"): CHECK, ("One-variable", "H18"): CHECK}),
    case("text_pass_cost", {("Pass-through margin", "B7"): "oops"}, {("Pass-through margin", "B11"): CHECK, ("Pass-through margin", "B12"): CHECK, ("Pass-through margin", "B13"): CHECK, ("Pass-through margin", "B14"): CHECK}),
    case("blank_pass_cost", {("Pass-through margin", "B7"): None}, {("Pass-through margin", "B11"): CHECK, ("Pass-through margin", "B13"): CHECK}),
    case("negative_pass_cost", {("Pass-through margin", "B7"): -.1}, {("Pass-through margin", "B11"): CHECK, ("Pass-through margin", "B14"): CHECK}),
    case("zero_retention", {("Inputs", "E23"): 0}, {("Inputs", "B18"): "OK", ("Pass-through margin", "B11"): "OK", ("Pass-through margin", "B12"): 0, ("Pass-through margin", "B13"): -.42, ("Pass-through margin", "B14"): NO_FINITE}),
    case("zero_node_cut", {("Inputs", "B5"): 0, ("Inputs", "B6"): 1}, {("Inputs", "B18"): "OK", ("Pass-through margin", "B12"): 0, ("Pass-through margin", "B14"): NO_FINITE}),
    case("zero_network_fees", {("Inputs", "B23"): 0, ("Inputs", "B24"): 0}, {("Inputs", "B18"): "OK", ("Cooperative", "I6"): UNDEFINED, ("Cooperative", "N9"): UNDEFINED, ("Ledger & limits", "B20"): UNDEFINED, ("Ledger & limits", "B26"): "OK", ("Ledger & limits", "B27"): UNDEFINED, ("One-variable", "H7"): UNDEFINED}),
    case("zero_network_stake", {("Inputs", "C23"): 0, ("Inputs", "C24"): 0}, {("Inputs", "B18"): "OK", ("Cooperative", "J6"): UNDEFINED, ("Cooperative", "N9"): UNDEFINED, ("Ledger & limits", "B21"): UNDEFINED, ("Ledger & limits", "B27"): UNDEFINED, ("One-variable", "H18"): UNDEFINED}),
]


def main():
    binary = shutil.which("libreoffice")
    if binary is None:
        raise SystemExit("LibreOffice is required for this optional recalculation check")
    if not OUTPUT.is_file():
        raise SystemExit(f"Build workbook first: {OUTPUT}")
    original_hash = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    with TemporaryDirectory(prefix=".recalculate-", dir=HERE) as temp:
        base = Path(temp)
        outdir = base / "recalculated"
        outdir.mkdir()
        files = []
        for name, edits, _ in CASES:
            book = load_workbook(OUTPUT)
            for (sheet, address), value in edits.items():
                book[sheet][address] = value
            path = base / f"{name}.xlsx"
            book.save(path)
            files.append(path)
        profile = (base / "lo-profile").resolve().as_uri()
        command = [binary, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), *map(str, files)]
        result = subprocess.run(command, capture_output=True, text=True, timeout=90)
        if result.returncode:
            raise AssertionError(f"LibreOffice conversion failed: {result.stdout}\n{result.stderr}")
        for name, edits, expected in CASES:
            path = outdir / f"{name}.xlsx"
            assert path.is_file(), f"Missing recalculated workbook: {name}\n{result.stdout}\n{result.stderr}"
            values = load_workbook(path, data_only=True)
            for (sheet, address), target in expected.items():
                expect(values[sheet][address].value, target, f"{name} {sheet}!{address}")
            for sheet in values.sheetnames:
                for row in values[sheet]:
                    for cell in row:
                        if (sheet, cell.coordinate) not in edits:
                            assert cell.data_type != "e", f"Unexpected formula error: {name} {sheet}!{cell.coordinate}={cell.value}"
    assert hashlib.sha256(OUTPUT.read_bytes()).hexdigest() == original_hash, "Deliverable changed during regression"
    print(f"{len(CASES)} LibreOffice recalculation cases: PASS; original workbook unchanged")


if __name__ == "__main__":
    main()
